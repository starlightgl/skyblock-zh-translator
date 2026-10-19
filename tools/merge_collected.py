"""Merge text files collected by players into one workbook for translation.

Usage: python tools/merge_collected.py <文件或文件夹> [...] [-o 输出.xlsx]
       (默认输出 data/收集汇总.xlsx)

Input: config/npc_translator_text.tsv from any number of players (the older
npc_translator_menu_text.tsv is accepted too). Each file counts as one player.

Output sheets:
  说明              files read and counts per category
  收集-<类别>       text not yet in the term table, most widely seen first; fill 校对译名
  收集-已在术语表   text the term table already covers, for review only

tools/import_terms.py reads the 收集-* sheets (only rows whose 校对译名 is filled), so the
workbook can be passed to it as a supplement after translating.
"""
import argparse
import glob
import os
import re
import sys
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABLE = os.path.join(ROOT, "data/SkyBlock术语表.xlsx")
ORDER = ["物品名称", "物品描述", "菜单标题", "侧边栏", "Tab列表", "屏幕标题", "屏幕副标题", "动作栏",
         "Boss血条", "聊天", "名称与悬浮文字", "其他模组提示", "价格表格"]
COLUMNS = ["英文", "类型", "出现人数", "出现位置", "原文示例", "当前译名", "校对译名", "备注"]
INVALID_SHEET = re.compile(r"[\[\]:*?/\\]")


def input_files(paths):
    out = []
    for p in paths:
        if os.path.isdir(p):
            out += sorted(glob.glob(os.path.join(p, "**", "*.tsv"), recursive=True))
        elif os.path.isfile(p):
            out.append(p)
        else:
            print("not found:", p, file=sys.stderr)
    return out


def read_file(path):
    """Yields (category, place, template, example, current)."""
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        header = f.readline().rstrip("\r\n").split("\t")
        old = header and header[0] == "菜单"
        if not old and header and header[0] != "类别":
            f.seek(0)
        for line in f:
            p = line.rstrip("\r\n").split("\t")
            if len(p) < 5:
                continue
            if old:
                cat = "物品名称" if p[1] == "名称" else "物品描述"
                yield cat, p[0], p[2], p[3], p[4]
            else:
                yield p[0], p[1], p[2], p[3], p[4]


def table_terms():
    """English → current translation of everything the term table covers."""
    known = {}
    if not os.path.exists(TABLE):
        return known
    wb = load_workbook(TABLE, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or "英文" not in header:
            continue
        i_en = header.index("英文")
        i_zh = [header.index(k) for k in ("校对译名", "当前译名") if k in header]
        for r in rows:
            en = r[i_en] if i_en < len(r) else None
            if not isinstance(en, str) or not en.strip():
                continue
            zh = next((r[i] for i in i_zh if i < len(r) and isinstance(r[i], str) and r[i].strip()), None)
            if zh:
                known.setdefault(en.strip(), zh.strip())
    return known


def sheet_name(cat):
    return INVALID_SHEET.sub("_", "收集-" + cat)[:31]


def write_sheet(ws, rows, with_category=False):
    cols = (["类别"] if with_category else []) + COLUMNS
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in rows:
        ws.append(([r["cat"]] if with_category else []) + [r[k] for k in COLUMNS])
        for c in ws[ws.max_row]:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"  # text such as "=== Rewards ===" is not a formula
    widths = {"类别": 12, "英文": 60, "类型": 6, "出现人数": 8, "出现位置": 28, "原文示例": 50,
              "当前译名": 40, "校对译名": 40, "备注": 20}
    for i, name in enumerate(cols):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = widths.get(name, 14)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "B2" if not with_category else "C2"
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("-o", "--output", default=os.path.join(ROOT, "data/收集汇总.xlsx"))
    args = ap.parse_args()

    files = input_files(args.inputs)
    if not files:
        sys.exit("no input files")

    entries = {}
    players = defaultdict(set)
    places = defaultdict(lambda: defaultdict(int))
    for fi, path in enumerate(files):
        for cat, place, template, example, current in read_file(path):
            template = template.strip()
            if not template:
                continue
            key = (cat, template)
            e = entries.get(key)
            if e is None:
                e = entries[key] = {"cat": cat, "英文": template, "原文示例": example, "当前译名": current}
            elif not e["当前译名"] and current:
                e["当前译名"] = current
            players[key].add(fi)
            if place:
                places[key][place] += 1

    known = table_terms()
    todo = defaultdict(list)
    covered = []
    for key, e in entries.items():
        cat = e["cat"]
        e["类型"] = "名称" if cat == "物品名称" else "描述"
        e["出现人数"] = len(players[key])
        top = sorted(places[key].items(), key=lambda kv: -kv[1])[:3]
        e["出现位置"] = " / ".join(p for p, _ in top)
        e["校对译名"] = ""
        e["备注"] = ""
        if e["英文"] in known:
            e["当前译名"] = known[e["英文"]]
            covered.append(e)
        else:
            todo[cat].append(e)

    def sort_key(e):
        return (-e["出现人数"], e["英文"].lower())

    wb = Workbook()
    info = wb.active
    info.title = "说明"
    info.append(["收集汇总"])
    info["A1"].font = Font(bold=True, size=14)
    info.append([f"读取文件 {len(files)} 个（每个文件计为一名玩家），共 {len(entries)} 条不重复原文。"])
    info.append(["「收集-<类别>」为术语表尚未覆盖的原文，按出现人数排序；填写「校对译名」后可用 tools/import_terms.py 导入。"])
    info.append(["「当前译名」是收集时游戏里显示的译文（可能来自社区语料或 AI），仅供参考，不会被导入。"])
    info.append(["「{0}」「{1}」为数字占位，译文中需保留；「{玩家}」为玩家本人的游戏名。"])
    info.append([])
    info.append(["类别", "待翻译", "已在术语表"])
    for c in info[7]:
        c.font = Font(bold=True)
    cats = [c for c in ORDER if c in todo] + sorted(c for c in todo if c not in ORDER)
    covered_by = defaultdict(int)
    for e in covered:
        covered_by[e["cat"]] += 1
    for c in cats + sorted(c for c in covered_by if c not in todo):
        info.append([c, len(todo.get(c, [])), covered_by.get(c, 0)])
    info.append([])
    info.append(["文件"])
    for p in files:
        info.append([os.path.basename(os.path.dirname(p)) + "/" + os.path.basename(p)])
    info.column_dimensions["A"].width = 40
    info.column_dimensions["B"].width = 10
    info.column_dimensions["C"].width = 12

    for c in cats:
        write_sheet(wb.create_sheet(sheet_name(c)), sorted(todo[c], key=sort_key))
    if covered:
        write_sheet(wb.create_sheet("收集-已在术语表"), sorted(covered, key=lambda e: (e["cat"], e["英文"].lower())), True)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    wb.save(args.output)
    print(f"{len(files)} files, {len(entries)} entries, {sum(len(v) for v in todo.values())} to translate -> {args.output}")
    for c in cats:
        print(f"  {c}: {len(todo[c])}")


if __name__ == "__main__":
    main()
